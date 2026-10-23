import numpy as np
import pytest

import analisis as an


def test_validacion_de_ventas():
    assert an.validar_ventas(an.cargar("a_ventas.csv")) == {"filas": 109, "duplicadas": 1, "ausentes": 5}


def test_totales_por_region():
    t = an.ventas_por_region(an.cargar("a_ventas.csv")).set_index("region")
    assert t.loc["Norte", "total_eur"] == pytest.approx(1911061.69, abs=0.01)
    assert t.loc["Centro", "total_eur"] == pytest.approx(2539781.72, abs=0.01)
    assert t.loc["Sur", "total_eur"] == pytest.approx(1619001.97, abs=0.01)
    assert t["total_eur"].sum() == pytest.approx(6069845.38, abs=0.01)
    assert t["meses"].tolist() == [34, 35, 34]


def test_fechas_mixtas():
    b = an.cargar("b_absentismo.csv")
    assert int(b["fecha"].str.contains("/").sum()) == 21
    assert an.parsear_fechas(b["fecha"]).isna().sum() == 0
    assert int(b["dias_ausencia"].sum()) == 814


def test_hueco_del_sensor():
    assert an.huecos(an.cargar("d_ambiental.csv")["temperatura_C"]) == [(601, 613)]
    assert an.huecos([1, np.nan, np.nan, 4, np.nan]) == [(2, 3), (5, 5)]


def test_mantenimiento_por_tipo():
    c = an.cargar("c_mantenimiento.csv")
    g = c.groupby("tipo")
    assert g.size()["correctivo"] == 50 and g.size()["preventivo"] == 130
    assert g["horas_parada"].sum()["correctivo"] == pytest.approx(474.4, abs=0.05)
    assert int(c["coste_eur"].isna().sum()) == 3


def test_espectro_de_un_seno_puro():
    fs, n = 200.0, 4096
    t = np.arange(n) / fs
    f, A = an.espectro(2.0 * np.sin(2 * np.pi * 12.5 * t), fs)   # 12,5 Hz cae justo en un bin (256 · fs/n)
    assert f[np.argmax(A)] == pytest.approx(12.5, abs=1e-9)
    assert A.max() == pytest.approx(2.0, rel=0.01)


def test_espectro_del_caso_e_modo_principal_cerca_de_12_5_hz():
    e = an.cargar("e_vibracion.csv")
    f, A = an.espectro(e["aceleracion_m_s2"], 200.0)
    assert abs(f[np.argmax(A)] - 12.5) < 0.25      # pico ancho (ζ = 1,5 %): error de décimas de hercio


def test_excel_de_ida_y_vuelta(tmp_path):
    import excel
    ruta = excel.escribir_libro(tmp_path / "v.xlsx", an.cargar("a_ventas.csv"))
    assert pd_sheetnames(ruta) == ["resumen", "validacion"]
    assert len(excel.leer_hoja(ruta, "resumen")) == 3
    assert excel.leer_hoja(ruta, "validacion")["ausentes"].tolist() == [5]


def pd_sheetnames(ruta):
    import openpyxl
    return openpyxl.load_workbook(ruta).sheetnames
